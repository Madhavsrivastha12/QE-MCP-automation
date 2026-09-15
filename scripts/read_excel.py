#!/usr/bin/env python3
"""Read Excel file and print contents."""
import openpyxl

wb = openpyxl.load_workbook('C:/Users/TI/Downloads/Usage Empire - Feature 1 - Phase 1_526925 _ Custom forecast upload screen.xlsx')
print('Sheets:', wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n{"="*80}')
    print(f'Sheet: {sheet_name}')
    print(f'{"="*80}')

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 30:  # Print first 30 rows
            print(f'Row {i:2d}: {row}')
        else:
            break
