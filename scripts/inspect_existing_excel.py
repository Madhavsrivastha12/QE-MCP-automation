#!/usr/bin/env python3
"""
Inspect existing generated Excel files
"""

import pandas as pd
from openpyxl import load_workbook
from pathlib import Path

def inspect_existing_excel(file_path, file_name):
    """Inspect existing Excel file"""
    print(f"\n{'='*80}")
    print(f"INSPECTING: {file_name}")
    print(f"PATH: {file_path}")
    print(f"{'='*80}\n")

    # Load workbook
    wb = load_workbook(file_path)

    print(f"Total Sheets: {len(wb.sheetnames)}")
    print(f"Sheet Names: {wb.sheetnames}\n")

    # Inspect each sheet
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        print(f"--- SHEET: {sheet_name} ---")
        print(f"Max Row: {ws.max_row}")
        print(f"Max Column: {ws.max_column}\n")

        # Read headers (row 1)
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(1, col_idx)
            headers.append(str(cell.value).strip() if cell.value else '')

        print("COLUMN HEADERS:")
        for idx, header in enumerate(headers, 1):
            if header:
                print(f"  {idx}. {header}")
        print()

        # Show row count
        data_rows = ws.max_row - 1  # Exclude header
        print(f"DATA ROWS: {data_rows}\n")

        print(f"{'-'*80}\n")

def main():
    # Existing generated files
    files = [
        (r"c:\UE_Automation\QE_MCP_automation\outputs\643243\test-scenario-mapping.xlsx", "PBI 643243 - Test Scenarios"),
        (r"c:\UE_Automation\QE_MCP_automation\outputs\643243\Test_Cases_643243.xlsx", "PBI 643243 - Test Cases"),
        (r"c:\UE_Automation\QE_MCP_automation\outputs\645352\Test-Scenarios-Mapped-to-AC.xlsx", "PBI 645352 - Test Scenarios"),
        (r"c:\UE_Automation\QE_MCP_automation\outputs\645352\Test_Cases.xlsx", "PBI 645352 - Test Cases"),
    ]

    for file_path, file_name in files:
        if Path(file_path).exists():
            inspect_existing_excel(file_path, file_name)
        else:
            print(f"ERROR: File not found: {file_path}\n")

if __name__ == '__main__':
    main()